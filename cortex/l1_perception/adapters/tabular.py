"""CSV/XLSX upload adapter (investor lists, CRM exports, opportunity lists). Rows become raw items keyed by header."""

from __future__ import annotations

import csv
import io
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from typing import TYPE_CHECKING, Any

from cortex.l1_perception.adapters.base import FetchContext, register
from cortex.l1_perception.models import RawItem

if TYPE_CHECKING:
    from cortex.l1_perception.registry import JoinSheet, SourceConfig


class SheetNotFound(ValueError):
    """The requested worksheet isn't in the workbook (the upload is rejected with the sheets that are)."""

    def __init__(self, sheet: str, available: list[str]):
        self.sheet, self.available = sheet, available
        super().__init__(f"sheet '{sheet}' not found; available: {available}")


def _is_xlsx(filename: str) -> bool:
    return filename.lower().endswith((".xlsx", ".xlsm"))


def sheet_names(data: bytes, filename: str) -> list[str]:
    """Worksheet names of an XLSX upload ([] for CSV)."""
    if not _is_xlsx(filename):
        return []
    from openpyxl import load_workbook

    return list(load_workbook(io.BytesIO(data), read_only=True).sheetnames)


def _uncached_formulas(data: bytes, title: str, cached: list[tuple[Any, ...]]) -> set[tuple[int, int]]:
    """(row, col) of formula cells with no cached value. Formulas are never evaluated: such a cell fails its row."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=False)[title]
    out: set[tuple[int, int]] = set()
    for r, row in enumerate(ws.iter_rows(values_only=True)):
        for c, v in enumerate(row):
            if isinstance(v, str) and v.startswith("=") and r < len(cached) and c < len(cached[r]):
                if cached[r][c] is None:
                    out.add((r, c))
    return out


def read_rows(data: bytes, filename: str, sheet: str | None = None, *, with_meta: bool = False) -> list[dict[str, Any]]:
    """Rows keyed by header. XLSX: ``sheet`` names the worksheet (None = the first, as before). With ``with_meta``
    each XLSX row also carries its sheet row number in ``_row`` and, when a formula cell has no cached value, the
    columns in ``_uncached_formula`` (the ingestion path fails that row; other callers see plain rows)."""
    if _is_xlsx(filename):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise SheetNotFound(sheet, list(wb.sheetnames))
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        uncached = _uncached_formulas(data, ws.title, rows) if with_meta else set()
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        out = []
        for r, values in enumerate(rows[1:], start=1):
            missing = [header[c] for (rr, c) in sorted(uncached) if rr == r and c < len(header)] if with_meta else []
            if not missing and all(v in (None, "") for v in values):
                continue
            rec: dict[str, Any] = {
                header[i]: (v.isoformat() if isinstance(v, datetime) else v)
                for i, v in enumerate(values)
                if i < len(header)
            }
            if with_meta:
                rec["_row"] = r + 1
            if missing:
                rec["_uncached_formula"] = missing
            out.append(rec)
        return out
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    return [
        {(k or "").strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items()}
        for row in csv.DictReader(io.StringIO(text), dialect=dialect)
        if any((v or "").strip() for v in row.values() if isinstance(v, str))
    ]


def _key(v: Any) -> str:
    return str(v).strip().casefold() if v not in (None, "") else ""


def join_indexes(cfg: SourceConfig, data: bytes, filename: str) -> list[tuple[JoinSheet, dict[str, dict[str, Any]]]]:
    """({key: row} of each ``join_sheets`` sheet present in the workbook; a missing sheet is skipped."""
    if not cfg.join_sheets or not _is_xlsx(filename):
        return []
    available = set(sheet_names(data, filename))
    return [
        (j, {_key(r.get(j.key)): r for r in read_rows(data, filename, j.sheet) if _key(r.get(j.key))})
        for j in cfg.join_sheets
        if j.sheet in available
    ]


def apply_joins(row: dict[str, Any], joins: list[tuple[JoinSheet, dict[str, dict[str, Any]]]]) -> bool:
    """Merge each joined sheet's matching row into ``row`` under its prefix. True when any sheet matched."""
    hit = False
    for j, index in joins:
        match = index.get(_key(row.get(j.main_key)))
        if match:
            row.update({f"{j.prefix}{h}": v for h, v in match.items()})
            hit = True
    return hit


class TabularAdapter:
    name = "tabular"
    version = "tabular-1.0"

    async def fetch(self, cfg: SourceConfig, ctx: FetchContext) -> AsyncIterator[RawItem]:
        if ctx.upload is None:
            raise ValueError("tabular sources ingest uploaded files; use POST /v1/sources/{id}/upload")
        sheet = ctx.sheet or cfg.sheet
        name = ctx.upload_name or "upload.csv"
        joins = join_indexes(cfg, ctx.upload, name)
        for n, row in enumerate(read_rows(ctx.upload, name, sheet, with_meta=True)):
            if n >= ctx.max_items:
                break
            apply_joins(row, joins)
            r = row.pop("_row", n + 2)
            where = f"sheet={sheet}&row={r}" if sheet else f"row={r}"
            yield RawItem(
                payload={"_row": r, "_file": ctx.upload_name, **({"_sheet": sheet} if sheet else {}), **row},
                url=f"upload://{ctx.upload_name}#{where}",
                fetched_at=datetime.now(UTC),
            )


register(TabularAdapter())
