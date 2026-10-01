"""Compliance report (screen 16): one workbook an auditor can hand over — hash-chain verification, approval
decisions, external releases with their approvers, blocked releases, citation-check outcomes, legal holds,
retention runs and every settings / role change in the period. Every row is read from the database as-is."""

from __future__ import annotations

import io
from typing import Any

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from cortex.l7_governance import audit_service
from platform_core.config import get_settings


async def build_report(s: AsyncSession, date_from: str, date_to: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    org = get_settings().org_id
    p = {"org": org, "a": date_from, "b": date_to}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    chain = await audit_service.verify(s)
    ws.append([f"Capital Cortex compliance report {date_from} to {date_to}"])
    ws["A1"].font = Font(bold=True, size=13)

    async def q(sql: str) -> list[dict[str, Any]]:
        return [dict(r) for r in (await s.execute(text(sql), p)).mappings().all()]

    counts = {
        "Audit chain": "intact" if chain["ok"] else f"BROKEN at #{chain['first_broken_seq']}: {chain['reason']}",
        "Audit records checked": chain["checked"],
        "Approval decisions": (
            await q(
                "SELECT count(*) AS n FROM approval_decision WHERE org_id = :org AND created_at::date BETWEEN :a AND :b"
            )
        )[0]["n"],
        "External releases (sent)": (
            await q(
                "SELECT count(*) AS n FROM outbox WHERE org_id = :org AND status = 'sent' AND sent_at::date BETWEEN :a AND :b"
            )
        )[0]["n"],
        "Blocked releases": (
            await q(
                "SELECT count(*) AS n FROM audit_log WHERE org_id = :org AND action = 'outbox.blocked' AND ts::date BETWEEN :a AND :b"
            )
        )[0]["n"],
        "Invalidated approvals": (
            await q(
                "SELECT count(*) AS n FROM approval WHERE org_id = :org AND decision = 'invalidated' AND updated_at::date BETWEEN :a AND :b"
            )
        )[0]["n"],
        "Active legal holds": (
            await q("SELECT count(*) AS n FROM legal_hold WHERE org_id = :org AND released_at IS NULL")
        )[0]["n"],
    }
    for k, v in counts.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34

    def sheet(name: str, rows: list[dict[str, Any]]) -> None:
        sh = wb.create_sheet(name)
        if not rows:
            sh.append(["No records in the period."])
            return
        cols = list(rows[0])
        sh.append(cols)
        for c in sh[1]:
            c.font = Font(bold=True)
        for r in rows:
            sh.append([str(r[c]) if r[c] is not None else "" for c in cols])

    sheet("Approval decisions", await q(
        "SELECT d.created_at, a.subject_type, a.subject_id, d.approver_username, d.roles, d.mfa, d.decision, d.comment, d.content_hash "
        "FROM approval_decision d JOIN approval a ON a.id = d.approval_id WHERE d.org_id = :org AND d.created_at::date BETWEEN :a AND :b ORDER BY d.created_at"))  # fmt: skip
    sheet("External releases", await q(
        "SELECT o.sent_at, o.channel, o.kind, o.recipient, o.content_hash, a.requested_by, "
        "(SELECT string_agg(d.approver_username, ', ') FROM approval_decision d WHERE d.approval_id = a.id AND d.decision = 'approved') AS approvers "
        "FROM outbox o LEFT JOIN approval a ON a.id = o.approval_id WHERE o.org_id = :org AND o.status = 'sent' AND o.sent_at::date BETWEEN :a AND :b ORDER BY o.sent_at"))  # fmt: skip
    sheet("Blocked releases", await q(
        "SELECT ts, actor, target, meta ->> 'reason' AS reason FROM audit_log WHERE org_id = :org AND action = 'outbox.blocked' AND ts::date BETWEEN :a AND :b ORDER BY ts"))  # fmt: skip
    sheet("Citation checks", await q(
        "SELECT created_at, id AS recommendation_id, stance, method, citation_report ->> 'status' AS status, citation_report ->> 'passed' AS passed, "
        "citation_report ->> 'rejected' AS stripped FROM recommendation WHERE org_id = :org AND created_at::date BETWEEN :a AND :b ORDER BY created_at"))  # fmt: skip
    sheet(
        "Legal holds",
        await q(
            "SELECT created_at, target_table, target_id, scope, reason, created_by, released_at, released_by FROM legal_hold WHERE org_id = :org ORDER BY created_at"
        ),
    )
    sheet("Retention runs", await q("SELECT created_at, policy, target_table, action, dry_run, rows_affected, rows_held, run_by FROM retention_run "
                                    "WHERE org_id = :org AND created_at::date BETWEEN :a AND :b ORDER BY created_at"))  # fmt: skip
    sheet("Settings and roles", await q(
        "SELECT ts, actor, action, target, meta FROM audit_log WHERE org_id = :org AND (action LIKE 'setting.%' OR action LIKE 'admin.user.%' "
        "OR action LIKE 'legal_hold.%' OR action = 'proposal.gap_waived') AND ts::date BETWEEN :a AND :b ORDER BY ts"))  # fmt: skip
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
