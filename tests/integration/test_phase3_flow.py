"""Phase 3 definition of done on the real Postgres (AGE + pgvector) and real OPA/Rego:
a VC package with a live-formula financial model and evidence appendix (gaps block approval until waived), a DD
package from approved documents with an approval-gated share link, a board pack approved by an Executive and
distributed through the outbox, the Copilot (grounded answer and honest refusal), retention with legal hold,
admin settings, and the compliance report."""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
import zipfile
from datetime import UTC, date, datetime, timedelta

import fakeredis.aioredis
import httpx
import pytest
from sqlalchemy import text

from cortex.l7_governance import outbox as outbox_mod
from cortex.l8_actuation.api.app import create_app
from platform_core import objectstore
from platform_core.auth import oidc
from platform_core.bus import streams
from platform_core.db import session_scope
from platform_core.policy import opa
from tests.integration.test_phase2_flow import _seed_high_opportunity

pytestmark = pytest.mark.integration
KEY = "integration-approval-signing-key-0123456789"


@pytest.fixture
async def env(engine, verifier, opa_url, monkeypatch):
    monkeypatch.setenv("APPROVAL_SIGNING_KEY", KEY)
    store: dict[tuple[str, str], bytes] = {}

    async def put(bucket, key, data, content_type="application/octet-stream"):
        store[(bucket, key)] = data
        return f"s3://{bucket}/{key}"

    async def get(bucket, key):
        return store[(bucket, key)]

    monkeypatch.setattr(objectstore, "put_bytes", put)
    monkeypatch.setattr(objectstore, "get_bytes", get)
    sent: list = []
    monkeypatch.setattr(outbox_mod, "_smtp_send", lambda msg: sent.append(msg) or "smtp://test")
    oidc.set_verifier(verifier)
    opa.set_opa(opa.OPAClient(opa_url))
    streams.set_bus(streams.Bus(fakeredis.aioredis.FakeRedis()))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app()), base_url="http://t", timeout=120
    ) as c:
        yield c, sent, store
    oidc.set_verifier(None)
    opa.set_opa(None)
    streams.set_bus(None)


async def _enrich_profile_and_financials() -> None:
    src = json.dumps({"kind": "manual_entry", "entered_by": "integration"})
    async with session_scope() as s:
        me = (
            (await s.execute(text("SELECT id, profile FROM organization WHERE kind = 'self' AND NOT is_demo LIMIT 1")))
            .mappings()
            .one()
        )
        prof = {**me["profile"], "team": [{"name": "Ada Orlo", "role": "CEO", "bio": "Former grid engineer"}],
                "traction": ["3 utility pilots signed"], "use_of_funds": [{"item": "Pilot deployments", "share_pct": 60}],
                "budget_lines": [{"category": "Staff", "item": "Engineers", "amount": 180000}, {"category": "Equipment", "item": "Test rig", "amount": 40000}]}  # fmt: skip
        await s.execute(
            text("UPDATE organization SET profile = CAST(:p AS jsonb) WHERE id = :id"),
            {"p": json.dumps(prof), "id": me["id"]},
        )
        await s.execute(text("DELETE FROM financial_snapshot WHERE NOT is_demo"))
        for m in range(6):
            period = (date.today().replace(day=1) - timedelta(days=31 * (6 - m))).replace(day=1)
            await s.execute(
                text("INSERT INTO financial_snapshot (period, cash, revenue, opex, net_burn, currency, source_ref) VALUES "
                     "(:p, :c, :r, :o, :b, 'USD', CAST(:src AS jsonb)) ON CONFLICT (org_id, period) DO NOTHING"),
                {"p": period, "c": 1_200_000 - m * 60_000, "r": 20_000 + m * 2_000, "o": 80_000, "b": 60_000 - m * 2_000, "src": src},
            )  # fmt: skip


async def test_phase3_definition_of_done(env, make_token):
    client, sent, store = env
    analyst = {"Authorization": f"Bearer {make_token(['analyst'], sub='analyst-3', username='ana', auth_age=5)}"}
    admin = {"Authorization": f"Bearer {make_token(['admin', 'approver'], sub='admin-3', username='adm', auth_age=5)}"}
    executive = {"Authorization": f"Bearer {make_token(['executive'], sub='exec-3', username='exe', auth_age=5)}"}
    ids = await _seed_high_opportunity()
    await _enrich_profile_and_financials()

    # ---------------------------------------------------------------- VC package (FR-05)
    r = await client.post(
        "/v1/proposals", headers=analyst, json={"opportunity_id": ids["opp"], "package_type": "vc_pitch"}
    )
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    p = (await client.get(f"/v1/proposals/{pid}", headers=analyst)).json()
    assert set(p["artefacts"]) == {
        "pitch_deck",
        "executive_summary",
        "investment_memo",
        "financial_model",
        "dd_checklist",
    }
    claims = [c for s in p["sections"] for c in s["claims"]]
    assert len(claims) >= 10 and all(c["evidence"] or c["basis"] for c in claims), "every claim is cited"
    assert any("Ada Orlo" in c["text"] for c in claims) and any("3 utility pilots signed" == c["text"] for c in claims)
    assert p["open_gaps"] > 0, "missing evidence is shown as gaps, never filled in"

    exports = {}
    for art, fmt in (
        ("pitch_deck", "pptx"),
        ("executive_summary", "docx"),
        ("investment_memo", "docx"),
        ("financial_model", "xlsx"),
        ("dd_checklist", "xlsx"),
    ):
        r = await client.get(f"/v1/proposals/{pid}/export", headers=analyst, params={"artefact": art, "fmt": fmt})
        assert r.status_code == 200, (art, r.text)
        assert hashlib.sha256(r.content).hexdigest() == r.headers["X-Content-SHA256"]
        exports[art] = r.content
    from docx import Document
    from openpyxl import load_workbook

    memo = Document(io.BytesIO(exports["investment_memo"]))
    body = "\n".join(x.text for x in memo.paragraphs)
    assert "[EVIDENCE REQUIRED:" in body and "Evidence appendix" in body and re.search(r"\[\d+\]", body)
    fm = load_workbook(io.BytesIO(exports["financial_model"]))
    assert str(fm["Projection"]["G2"].value).startswith("=") and str(fm["Runway"]["B1"].value).startswith(
        "=IFERROR(MATCH("
    )
    assert fm["Inputs"]["B3"].value is not None, "the model starts from the imported snapshots"
    pdf = await client.get(
        f"/v1/proposals/{pid}/export", headers=analyst, params={"artefact": "executive_summary", "fmt": "pdf"}
    )
    assert pdf.status_code in (200, 503)  # WeasyPrint is present in the API image; bare Windows lacks Pango

    # gaps block approval → an Admin waives them (audited) → submit → approve
    r = await client.post(f"/v1/proposals/{pid}/submit", headers=analyst)
    assert r.status_code == 409 and r.json()["title"] == "Evidence required"
    r = await client.post(
        f"/v1/proposals/{pid}/sections/team/gaps/x/waive", headers=analyst, json={"reason": "analysts can't waive"}
    )
    assert r.status_code == 403
    p = (await client.get(f"/v1/proposals/{pid}", headers=analyst)).json()
    for g in [g for g in p["gaps"] if not g["waived"]]:
        r = await client.post(f"/v1/proposals/{pid}/sections/{g['section']}/gaps/{g['id']}/waive", headers=admin,
                              json={"reason": "accepted for the integration test"})  # fmt: skip
        assert r.status_code == 200, r.text
    r = await client.post(f"/v1/proposals/{pid}/submit", headers=analyst)
    assert r.status_code == 200, r.text
    aid = r.json()["approval_id"]
    r = await client.post(f"/v1/approvals/{aid}/decision", headers=admin, json={"decision": "approved"})
    assert r.json()["status"] == "approved", r.text
    assert (await client.get(f"/v1/proposals/{pid}", headers=analyst)).json()["status"] == "approved"
    # editing after approval invalidates it
    sec = next(s for s in p["sections"] if s["claims"])
    r = await client.put(
        f"/v1/proposals/{pid}/sections/{sec['key']}", headers=analyst, json={"claims": sec["claims"][:1]}
    )
    assert r.status_code == 200 and r.json()["approvals_invalidated"] == 1
    assert (await client.get(f"/v1/proposals/{pid}", headers=analyst)).json()["status"] == "draft"
    bad = await client.put(f"/v1/proposals/{pid}/sections/{sec['key']}", headers=analyst,
                           json={"claims": [{"text": "We have 900 customers.", "kind": "fact", "evidence": [f"opportunity:{ids['opp']}"]}]})  # fmt: skip
    assert bad.status_code == 422, "a person's edit is held to the citation checker too"

    # ---------------------------------------------------------------- data room + DD package + share link (FR-06)
    docs = []
    for title, tags in (("Cap table", "cap_table"), ("Financial statements 2025", "financials")):
        r = await client.post("/v1/dataroom/documents", headers=analyst, files={"file": (f"{title}.pdf", f"%PDF {title}".encode())},
                              data={"title": title, "folder": "/finance", "dd_tags": tags, "classification": "confidential"})  # fmt: skip
        assert r.status_code == 201, r.text
        docs.append(r.json()["id"])
    v2 = await client.post("/v1/dataroom/documents", headers=analyst, files={"file": ("Cap table.pdf", b"%PDF v2")},
                           data={"title": "Cap table", "folder": "/finance", "dd_tags": "cap_table"})  # fmt: skip
    assert v2.json()["version"] == 2 and v2.json()["previous_version_id"] == docs[0]
    docs[0] = v2.json()["id"]
    r = await client.post("/v1/dataroom/packages", headers=analyst, json={"name": "Seed DD", "document_ids": docs})
    assert r.status_code == 422, "only approved documents can be packaged"
    r = await client.patch(f"/v1/dataroom/documents/{docs[0]}", headers=analyst, json={"approved_repo": True})
    assert r.status_code == 403, "approved-repo toggle is Admin/Legal only"
    for d in docs:
        assert (
            await client.patch(f"/v1/dataroom/documents/{d}", headers=admin, json={"approved_repo": True})
        ).status_code == 200
    cl = (await client.get("/v1/dataroom/checklist", headers=analyst)).json()
    assert {i["key"]: i["status"] for i in cl["items"]}["cap_table"] == "covered"
    r = await client.post("/v1/dataroom/packages", headers=analyst, json={"name": "Seed DD", "from_checklist": True})
    assert r.status_code == 201, r.text
    pkg = r.json()
    assert {f["title"] for f in pkg["manifest"]["files"]} == {"Cap table", "Financial statements 2025"}
    r = await client.post(
        f"/v1/dataroom/packages/{pkg['id']}/share",
        headers=analyst,
        json={"recipient": "partner@fund.example", "expires_in_days": 7},
    )
    share = r.json()
    r = await client.post(
        f"/v1/approvals/{share['approval_id']}/decision", headers=admin, json={"decision": "approved"}
    )
    assert r.json()["status"] == "approved"
    async with session_scope() as s:
        out = await outbox_mod.release(
            s, oidc.get_verifier().verify_sync(admin["Authorization"][7:]), share["outbox_id"]
        )
    assert out["status"] == "sent", out
    link = re.search(r"/v1/share/(\S+)", sent[-1].get_body().get_content()).group(1)
    z = await client.get(f"/v1/share/{link}")
    assert z.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(z.content))
    man = json.loads(zf.read("MANIFEST.json"))
    for f in man["files"]:
        assert hashlib.sha256(zf.read(f["path"])).hexdigest() == f["sha256"], "package checksums verify"
    assert (await client.get("/v1/share/" + "x" * 43)).status_code == 404
    async with session_scope() as s:
        await s.execute(
            text("UPDATE share_link SET expires_at = now() - interval '1 minute' WHERE id = :id"),
            {"id": share["share_link_id"]},
        )
    assert (await client.get(f"/v1/share/{link}")).status_code == 410
    log = (await client.get("/v1/dataroom/access-log", headers=analyst)).json()["items"]
    assert {"upload", "package", "share_download"} <= {x["action"] for x in log}

    # ---------------------------------------------------------------- board pack: Executive approves → distribution
    today = date.today()
    r = await client.post("/v1/board-reports", headers=analyst, json={"period_start": str(today - timedelta(days=90)), "period_end": str(today),
                                                                     "recipients": ["board@inspironics.net"]})  # fmt: skip
    assert r.status_code == 201, r.text
    bid = r.json()["id"]
    b = (await client.get(f"/v1/board-reports/{bid}", headers=executive)).json()
    assert {s_["key"] for s_ in b["content"]["sections"]} >= {
        "runway",
        "pipeline",
        "key_opportunities",
        "risks",
        "asks",
    }
    assert any(c for s_ in b["content"]["sections"] for c in s_["claims"])
    pptx = await client.get(f"/v1/board-reports/{bid}/export", headers=executive, params={"fmt": "pptx"})
    assert pptx.status_code == 200
    r = await client.post(f"/v1/board-reports/{bid}/submit", headers=analyst)
    baid = r.json()["approval_id"]
    r = await client.post(f"/v1/approvals/{baid}/decision", headers=executive, json={"decision": "approved"})
    assert r.status_code == 200 and r.json()["status"] == "approved", r.text
    dist = r.json()["follow_up"]["distribution"]
    assert len(dist) == 1 and dist[0]["approval"] == "approved"
    async with session_scope() as s:
        out = await outbox_mod.release(
            s, oidc.get_verifier().verify_sync(executive["Authorization"][7:]), dist[0]["outbox_id"]
        )
    assert out["status"] == "sent" and out["delivery"]["attachments"]
    b = (await client.get(f"/v1/board-reports/{bid}", headers=executive)).json()
    assert b["status"] == "distributed" and b["distribution"][0]["outbox_status"] == "sent"

    # ---------------------------------------------------------------- Copilot: grounded answer and honest refusal
    async def ask(q: str, opp: str | None = None) -> dict:
        events = []
        async with client.stream(
            "POST", "/v1/copilot/ask", headers=analyst, json={"question": q, "opportunity_id": opp}
        ) as st:
            async for line in st.aiter_lines():
                if line.startswith("data:"):
                    events.append(json.loads(line[5:]))
        return events[-1]

    done = await ask("Why is this ranked where it is?", ids["opp"])
    assert done["grounded"] and "ranked #" in done["answer"]
    done = await ask("What is the airspeed velocity of an unladen swallow?")
    assert not done["grounded"] and done["answer"] == "I don't have sourced evidence for that." and done["gaps"]

    # ---------------------------------------------------------------- retention with legal hold; admin settings; report
    r = await client.post(
        "/v1/legal-holds",
        headers=admin,
        json={"target_table": "proposal", "target_id": pid, "reason": "litigation hold"},
    )
    assert r.status_code == 201
    assert (await client.post("/v1/admin/retention/run", headers=admin, json={"dry_run": True})).status_code == 200
    from cortex.l7_governance import retention

    async with session_scope() as s:  # three years on: the draft is past its 1-year retention, but it is held
        run = await retention.run(s, "system:test", dry_run=False, now=datetime.now(UTC) + timedelta(days=3 * 365))
    prop = next(x for x in run["results"] if x["policy"] == "proposal_draft")
    assert prop["held"] >= 1
    assert (await client.get(f"/v1/proposals/{pid}", headers=analyst)).status_code == 200, "held rows are never deleted"
    r = await client.put("/v1/admin/policies", headers=admin, json={"governance": {"business_hours_only": False}})
    assert r.status_code == 200 and r.json()["version"] >= 1
    r = await client.put("/v1/admin/policies", headers=admin, json={"governance": {"rego": "allow := true"}})
    assert r.status_code == 422, "only whitelisted governance settings are editable"
    rep = await client.get(
        "/v1/audit/compliance-report", headers=admin, params={"from": str(today - timedelta(days=1)), "to": str(today)}
    )
    assert rep.status_code == 200
    wb = load_workbook(io.BytesIO(rep.content))
    assert {"Summary", "Approval decisions", "External releases", "Legal holds", "Settings and roles"} <= set(
        wb.sheetnames
    )
    assert wb["Summary"]["B2"].value == "intact"
    bad_id = await client.get("/v1/opportunities/0", headers=analyst)
    assert bad_id.status_code == 422 and bad_id.headers["content-type"].startswith("application/problem+json"), (
        "malformed ids are 422, not 500"
    )
    cost = (await client.get("/v1/dashboards/cost", headers=admin)).json()
    assert any(a["agent"] == "proposal" for a in cost["by_agent"])
    _ = (uuid, datetime, UTC, store)
