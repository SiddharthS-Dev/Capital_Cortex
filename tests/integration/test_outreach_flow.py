"""Capital outreach workbook (FR-04-OUT) through the real API on the real database."""

from __future__ import annotations

import json
from pathlib import Path

import fakeredis.aioredis
import httpx
import pytest
from sqlalchemy import text

from cortex.l2_representation.pipeline import process_signal
from cortex.l8_actuation.api.app import create_app
from platform_core.auth import oidc
from platform_core.bus import streams
from platform_core.db import session_scope
from platform_core.policy import opa
from platform_core.policy.opa import PolicyDecision

pytestmark = pytest.mark.integration

WORKBOOK = Path(__file__).resolve().parents[2] / "docs" / "Meris_Capital_Cortex_Outreach_Workbook_2026-10-06_9405.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class AllowOPA:
    async def evaluate(self, package, input_):
        return PolicyDecision(True, [])


@pytest.fixture
async def client(engine, verifier):
    oidc.set_verifier(verifier)
    opa.set_opa(AllowOPA())  # type: ignore[arg-type]
    streams.set_bus(streams.Bus(fakeredis.aioredis.FakeRedis()))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=create_app()), base_url="http://t") as c:
        yield c
    oidc.set_verifier(None)
    opa.set_opa(None)
    streams.set_bus(None)


async def drain_signals() -> int:
    bus = streams.get_bus()
    msgs = await bus.r.xrange("signals.raw")
    for _, fields in msgs:
        await process_signal(json.loads(fields[b"payload"])["signal_id"])
    await bus.r.delete("signals.raw")
    return len(msgs)


def _file() -> dict:
    return {"file": (WORKBOOK.name, WORKBOOK.read_bytes(), XLSX)}


async def _source_id(client, h, key: str) -> str:
    items = (await client.get("/v1/sources", headers=h)).json()["items"]
    return next(s["id"] for s in items if s["adapter_key"] == key)


async def test_workbook_inspect_upload_and_reupload(client, make_token):
    h = {"Authorization": f"Bearer {make_token(['admin', 'approver'], auth_age=5)}"}
    src = await _source_id(client, h, "capital_outreach")

    # a sheet that isn't there: 422 naming the ones that are, nothing written
    r = await client.post(f"/v1/sources/{src}/upload/inspect?sheet=Nope", headers=h, files=_file())
    assert r.status_code == 422 and "sheet 'Nope' not found; available: ['00_Read_Me'" in r.json()["detail"]
    r = await client.post(f"/v1/sources/{src}/upload?sheet=Nope", headers=h, files=_file())
    assert r.status_code == 422

    # dry run: the configured sheet, 27 headers matched, 52 would import, nothing stored
    r = await client.post(f"/v1/sources/{src}/upload/inspect", headers=h, files=_file())
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["sheet"] == "12_Meris_Import" and len(rep["headers"]) == 27 and rep["would_import"] == 52
    assert rep["unmapped_headers"] == [] and rep["required_missing"] == []
    async with session_scope() as s:
        n = await s.execute(text("SELECT count(*) FROM signal WHERE source_id = CAST(:s AS uuid)"), {"s": src})
        assert n.scalar() == 0  # the dry run stored nothing
        assert (
            await s.execute(text("SELECT count(*) FROM audit_log WHERE action = 'source.upload.inspect'"))
        ).scalar() == 1

    # upload: 52 new → 52 opportunities keyed by prospect id; no amount, currency or deadline anywhere
    r = await client.post(f"/v1/sources/{src}/upload", headers=h, files=_file())
    assert r.status_code == 200, r.text
    assert (r.json()["new"], r.json()["failed"], r.json()["status"]) == (52, 0, "succeeded")
    assert await drain_signals() == 52
    async with session_scope() as s:
        rows = (
            await s.execute(
                text(
                    "SELECT external_key, amount_min, amount_max, currency, deadline FROM opportunity "
                    "WHERE external_key LIKE 'capital_outreach:%'"
                )
            )
        ).all()
        assert len(rows) == 52 and {r[0] for r in rows} >= {"capital_outreach:CC-001", "capital_outreach:CC-052"}
        assert all(r[1] is None and r[2] is None and r[3] is None and r[4] is None for r in rows)
        prov = (
            await s.execute(
                text("SELECT source_ref FROM signal WHERE external_id = 'CC-002'"),
            )
        ).scalar_one()
        assert prov["url"].endswith("#sheet=12_Meris_Import&row=3")

    # the same workbook again: 52 duplicates, no new opportunities
    r = await client.post(f"/v1/sources/{src}/upload", headers=h, files=_file())
    assert (r.json()["new"], r.json()["duplicate"]) == (0, 52)
    async with session_scope() as s:
        n = (
            await s.execute(text("SELECT count(*) FROM opportunity WHERE external_key LIKE 'capital_outreach:%'"))
        ).scalar()
        assert n == 52


async def test_generic_upload_source_keeps_reading_the_first_sheet(client, make_token):
    h = {"Authorization": f"Bearer {make_token(['admin'], auth_age=5)}"}
    src = await _source_id(client, h, "csv_upload")
    r = await client.post(f"/v1/sources/{src}/upload", headers=h, files=_file())
    assert r.status_code == 200 and (r.json()["fetched"], r.json()["new"], r.json()["failed"]) == (16, 0, 16)
    r = await client.post(f"/v1/sources/{src}/upload?sheet=12_Meris_Import", headers=h, files=_file())
    assert (r.json()["fetched"], r.json()["failed"]) == (52, 52)  # no title header for the generic mapping


def _workbook_with(changes: dict[str, dict[str, object]]) -> bytes:
    """The real workbook with some 12_Meris_Import cells changed, built in memory (the file in docs/ is never re-saved)."""
    import io

    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK)
    ws = wb["12_Meris_Import"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        pid = row[0].value
        for col, value in changes.get(pid, {}).items():
            row[header.index(col)].value = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _profile(pid: str) -> dict:
    async with session_scope() as s:
        r = (
            (await s.execute(text("SELECT * FROM outreach_profile WHERE prospect_id = :p"), {"p": pid}))
            .mappings()
            .one()
        )
        return dict(r)


async def test_outreach_profiles_and_human_wins_on_reimport(client, make_token):
    import time

    h = {"Authorization": f"Bearer {make_token(['admin'], auth_age=5)}"}
    src = await _source_id(client, h, "capital_outreach")
    await client.post(f"/v1/sources/{src}/upload", headers=h, files=_file())  # new or already there
    await drain_signals()
    async with session_scope() as s:
        n = (await s.execute(text("SELECT count(*) FROM outreach_profile"))).scalar()
        assert n == 52
        events = (await s.execute(text("SELECT count(*) FROM outreach_status_event"))).scalar()
        assert events == 52  # one initial event per profile
    p = await _profile("CC-002")
    assert (p["route"], p["engagement_outlook"], p["analyst_priority"], p["outreach_status"]) == (
        "Contact now", "High", 90, "Not contacted",
    )  # fmt: skip
    assert p["phones"] == ["+1 423-281-0811"] and p["verified_on"].isoformat() == "2026-10-06"
    assert (p["source_ref"]["sheet"], p["source_ref"]["row"]) == ("12_Meris_Import", 3)
    assert p["source_ref"]["official_source_url"] == p["official_source_url"]
    assert p["source_ref"]["field_sources"]["attributes.route"] == "raw:route"

    # a person works the tracker
    async with session_scope() as s:
        await s.execute(
            text(
                "UPDATE outreach_profile SET outreach_status = 'Sent', first_sent_on = DATE '2026-10-08', "
                "notes = 'called them', status_set_by = 'ana' WHERE prospect_id = 'CC-002'"
            )
        )
    # a newer workbook: CC-002's research changed, and its sheet status says something else
    marker = f"Revised ask {time.time_ns()}"
    data = _workbook_with({"CC-002": {"next_action": marker, "status": "Declined", "route": "Watch next intake"}})
    r = await client.post(f"/v1/sources/{src}/upload", headers=h, files={"file": (WORKBOOK.name, data, XLSX)})
    assert (r.json()["new"], r.json()["duplicate"]) == (1, 51)
    assert await drain_signals() == 1
    p = await _profile("CC-002")
    assert p["next_action"] == marker and p["route"] == "Watch next intake" and p["import_status"] == "Declined"
    assert (p["outreach_status"], p["notes"], p["status_set_by"]) == ("Sent", "called them", "ana")  # human wins
    assert p["first_sent_on"].isoformat() == "2026-10-08"
    async with session_scope() as s:
        n = (
            await s.execute(text("SELECT count(*) FROM opportunity WHERE external_key LIKE 'capital_outreach:%'"))
        ).scalar()
        assert n == 52  # updated in place, no duplicate opportunity


async def test_outreach_status_events_are_append_only(engine):
    from sqlalchemy.exc import DBAPIError

    for stmt in ("UPDATE outreach_status_event SET reason = 'x'", "DELETE FROM outreach_status_event",
                 "TRUNCATE outreach_status_event"):  # fmt: skip
        with pytest.raises(DBAPIError, match="append-only"):
            async with session_scope() as s:
                await s.execute(text(stmt))


async def _ensure_workbook(client, h) -> None:
    src = await _source_id(client, h, "capital_outreach")
    await client.post(f"/v1/sources/{src}/upload", headers=h, files=_file())
    await drain_signals()


async def _opp_id(pid: str) -> str:
    return str((await _profile(pid))["opportunity_id"])


async def _opp(opp_id: str) -> dict:
    async with session_scope() as s:
        r = (
            (
                await s.execute(
                    text(
                        "SELECT pipeline_stage::text AS stage, status, owner_id FROM opportunity "
                        "WHERE id = CAST(:i AS uuid)"
                    ),
                    {"i": opp_id},
                )
            )
            .mappings()
            .one()
        )
        return dict(r)


async def _follow_ups(opp_id: str) -> list[dict]:
    async with session_scope() as s:
        rows = (
            (
                await s.execute(
                    text(
                        "SELECT title, due_at, owner_id, status FROM milestone WHERE opportunity_id = CAST(:i AS uuid) "
                        "AND kind = 'follow_up' AND source_ref->>'kind' = 'outreach_follow_up' ORDER BY due_at, status"
                    ),
                    {"i": opp_id},
                )
            )
            .mappings()
            .all()
        )
        return [dict(r) for r in rows]


async def test_status_machine_follow_ups_and_stage(client, make_token):
    from datetime import date

    from cortex.l3_memory import outreach_service as svc
    from platform_core.errors import Problem

    h = {"Authorization": f"Bearer {make_token(['admin'], auth_age=5)}"}
    await _ensure_workbook(client, h)
    actor = oidc.get_verifier().verify_sync(make_token(["analyst"], username="ana", auth_age=5))
    opp = await _opp_id("CC-003")

    # Sent: stage forward to engaged, +5/+12 follow-ups, unowned because the opportunity is unowned
    async with session_scope() as s:
        r = await svc.set_status(s, actor, opp, "Sent", first_sent_on=date(2026, 10, 8))
    assert r["stage_changed"] and r["pipeline_stage"] == "engaged" and len(r["follow_ups_created"]) == 2
    fus = await _follow_ups(opp)
    assert [f["due_at"].date().isoformat() for f in fus] == ["2026-10-13", "2026-10-20"]
    assert {f["owner_id"] for f in fus} == {None} and {f["status"] for f in fus} == {"open"}
    assert fus[0]["title"].startswith("Follow-up 1: ")
    async with session_scope() as s:
        await svc.set_status(s, actor, opp, "Sent")  # again: no second pair of follow-ups
    assert len(await _follow_ups(opp)) == 2

    # a reply cancels the open follow-ups; the stage stays engaged
    async with session_scope() as s:
        r = await svc.set_status(s, actor, opp, "Reply received", reason="replied by email")
    assert len(r["follow_ups_cancelled"]) == 2 and not r["stage_changed"]
    assert {f["status"] for f in await _follow_ups(opp)} == {"cancelled"}

    # forward only: Applied moves to submitted, then Prepared leaves it there
    async with session_scope() as s:
        await svc.set_status(s, actor, opp, "Applied")
        r = await svc.set_status(s, actor, opp, "Prepared")
    assert r["pipeline_stage"] == "submitted" and not r["stage_changed"]
    assert (await _opp(opp))["stage"] == "submitted"

    # Declined is allowed to go to lost, and closes the opportunity as lost
    async with session_scope() as s:
        await svc.set_status(s, actor, opp, "Declined")
    assert (await _opp(opp)) == {"stage": "lost", "status": "lost", "owner_id": None}

    async with session_scope() as s:
        events = (
            (
                await s.execute(
                    text(
                        "SELECT to_status FROM outreach_status_event WHERE opportunity_id = CAST(:i AS uuid) ORDER BY at"
                    ),
                    {"i": opp},
                )
            )
            .scalars()
            .all()
        )
        audited = (
            await s.execute(
                text("SELECT count(*) FROM audit_log WHERE action = 'outreach.status' AND target = :r"),
                {"r": f"opportunity:{opp}"},
            )
        ).scalar()
    # "Sent" twice: the second records an event and an audit row (the date may change), creates no follow-ups
    assert events == ["Not contacted", "Sent", "Sent", "Reply received", "Applied", "Prepared", "Declined"]
    assert audited == 6

    # follow-ups go to the opportunity owner when there is one
    owned = await _opp_id("CC-005")
    async with session_scope() as s:
        await s.execute(text("UPDATE opportunity SET owner_id = 'kumar' WHERE id = CAST(:i AS uuid)"), {"i": owned})
        await svc.set_status(s, actor, owned, "Sent", first_sent_on=date(2026, 10, 9))
    assert {f["owner_id"] for f in await _follow_ups(owned)} == {"kumar"}

    # Eligibility hold needs a linked gate or a recorded decision; it leaves the stage alone
    held = await _opp_id("CC-015")
    with pytest.raises(Problem, match="link an eligibility gate or record the eligibility decision"):
        async with session_scope() as s:
            await svc.set_status(s, actor, held, "Eligibility hold")
    async with session_scope() as s:
        r = await svc.set_status(s, actor, held, "Eligibility hold", eligibility_decision="G2: cap table pending")
    assert not r["stage_changed"] and (await _profile("CC-015"))["eligibility_decision"] == "G2: cap table pending"

    # Watchlist sets the opportunity status only
    watch = await _opp_id("CC-017")
    async with session_scope() as s:
        await svc.set_status(s, actor, watch, "Watchlist")
    assert (await _opp(watch))["status"] == "watchlist" and (await _opp(watch))["stage"] == "discovered"


async def test_contacts_owners_and_first_contact_draft(client, make_token):
    from cortex.l3_memory import outreach_service as svc

    h = {"Authorization": f"Bearer {make_token(['admin'], auth_age=5)}"}
    await _ensure_workbook(client, h)
    actor = oidc.get_verifier().verify_sync(make_token(["admin"], username="root", auth_age=5))
    created_kind = "SELECT count(*) FROM contact WHERE source_ref->>'kind' = 'outreach_contact_channel'"

    async with session_scope() as s:
        plan = await svc.import_contacts(s, actor, dry_run=True)
        assert (await s.execute(text(created_kind))).scalar() == 0  # the dry run wrote nothing
    assert len(plan["planned"]) == 31 and not plan["created"]
    assert {x["prospect_id"] for x in plan["skipped"]} >= {"CC-001", "CC-018", "CC-052"}
    async with session_scope() as s:
        done = await svc.import_contacts(s, actor, dry_run=False)
    assert len(done["created"]) == 31
    async with session_scope() as s:
        again = await svc.import_contacts(s, actor, dry_run=False)  # idempotent on the email
        rows = (
            (
                await s.execute(
                    text(
                        "SELECT name, role, emails, consent_basis, organization_id IS NOT NULL AS has_org FROM contact "
                        "WHERE source_ref->>'kind' = 'outreach_contact_channel' ORDER BY name"
                    )
                )
            )
            .mappings()
            .all()
        )
    assert not again["created"] and len(rows) == 31
    assert {r["consent_basis"] for r in rows} == {"public_professional"} and all(r["has_org"] for r in rows)
    assert next(r for r in rows if r["name"] == "Darrel Hugh")["emails"] == ["dhugh@ahla.com"]

    # owners: the map is empty, so nothing is assigned and every proposed name is reported
    async with session_scope() as s:
        o = await svc.apply_proposed_owners(s, actor, dry_run=False)
    assert o["applied"] == [] and sum(u["rows"] for u in o["unmapped"]) >= 40

    # first contact: a draft in the outbox, never sent
    opp = await _opp_id("CC-004")
    async with session_scope() as s:
        cid = (await s.execute(text("SELECT id FROM contact WHERE 'dhugh@ahla.com' = ANY(emails)"))).scalar()
        d = await svc.draft_first_contact(s, actor, opp, str(cid))
        ob = (
            (
                await s.execute(
                    text("SELECT status, recipient, payload, opportunity_id FROM outbox WHERE id = CAST(:i AS uuid)"),
                    {"i": d["id"]},
                )
            )
            .mappings()
            .one()
        )
    assert ob["status"] == "draft" and ob["recipient"] == "dhugh@ahla.com" and str(ob["opportunity_id"]) == opp
    assert ob["payload"]["body"].startswith("Dear Darrel Hugh,") and "Tailored first ask:" in ob["payload"]["body"]


async def test_outreach_alerts_next_action_and_staleness(client, make_token):
    from cortex.l8_actuation import alerts

    h = {"Authorization": f"Bearer {make_token(['admin'], auth_age=5)}"}
    await _ensure_workbook(client, h)
    async with session_scope() as s:
        await s.execute(text("UPDATE outreach_profile SET next_action_on = CURRENT_DATE WHERE prospect_id = 'CC-006'"))
        # research dates are workbook facts; moved back here only to exercise the refresh rules
        await s.execute(
            text("UPDATE outreach_profile SET verified_on = CURRENT_DATE - 100 WHERE prospect_id = 'CC-007'")
        )
        await s.execute(
            text(
                "UPDATE outreach_profile SET verified_on = CURRENT_DATE - 40 WHERE prospect_id IN ('CC-001', 'CC-030')"
            )
        )
        await alerts.seed_default_rules(s)
        await alerts.evaluate(s)
        rows = (
            await s.execute(
                text(
                    "SELECT r.name, p.prospect_id FROM alert a JOIN alert_rule r ON r.id = a.rule_id "
                    "JOIN outreach_profile p ON p.opportunity_id = a.subject_id WHERE r.name LIKE 'Outreach%'"
                )
            )
        ).all()
    got = {(n, p) for n, p in rows}
    assert ("Outreach next action due", "CC-006") in got
    assert ("Outreach sources stale", "CC-007") in got
    # a priority row (CC-001, priority 100) is stale after 30 days; a low-priority row (CC-030) only after 90
    assert ("Outreach research stale (priority rows)", "CC-001") in got
    assert not {g for g in got if g[1] == "CC-030"}


# last in this module: it drops and recreates the outreach tables
def test_migration_0007_downgrades_and_upgrades(pg_url):
    import os
    import subprocess
    import sys

    root = Path(__file__).resolve().parents[2]
    env = {**os.environ, "DATABASE_URL": pg_url}
    for target in ("0006_fx_rates", "head"):
        subprocess.run([sys.executable, "-m", "alembic", "downgrade" if target != "head" else "upgrade", target],
                       cwd=root, env=env, check=True)  # fmt: skip
