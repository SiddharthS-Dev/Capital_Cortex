"""Phase 3: asset generator (markers, appendix, gaps, live formulas), proposal composition, compliance rules,
retention durations, settings validation, extension points, Graph adapters."""

import io
from datetime import UTC, datetime

import httpx
import pytest
import respx

from cortex.l6_agency.tool_registry import ToolResult
from cortex.l7_governance import compliance_check
from cortex.l7_governance.retention import duration, policies
from cortex.l7_governance.settings_service import validate
from cortex.l8_actuation import asset_generator as ag
from cortex.l8_actuation.proposal_builder import compose, gap, packages
from platform_core.errors import Problem

SECTIONS = [
    {"key": "opportunity", "title": "The opportunity", "claims": [
        {"text": "The stated amount is USD 100,000 to 400,000.", "kind": "fact", "evidence": ["opportunity:1"], "basis": []},
        {"text": "The deadline is 2026-12-15.", "kind": "fact", "evidence": ["tool:r:opportunity", "opportunity:1"], "basis": []}],
     "gaps": [gap("opportunity", "no signal linked")]},
    {"key": "team", "title": "Team", "claims": [], "gaps": [gap("team", "team members (name, role, bio)")]},
]  # fmt: skip
LABELS = {
    "opportunity:1": ("opportunity: Clean energy grant", "grants_gov · https://x"),
    "tool:r:opportunity": ("opportunity: deterministic tool result", "computed"),
}


def model(waivers=None):
    return ag.build_model(title="Executive summary", subtitle="v1", artefact="executive_summary", sections=SECTIONS,
                          ref_labels=LABELS, waivers=waivers, meta={"version": 1, "content_hash": "abc", "mode": "deterministic"},
                          disclaimer="Not an offer.")  # fmt: skip


def test_citation_numbering_and_appendix():
    m = model()
    assert [i.marks for i in m.sections[0].items] == [[1], [2, 1]]
    assert [(a.n, a.ref) for a in m.appendix] == [(1, "opportunity:1"), (2, "tool:r:opportunity")]
    assert m.gap_count == 2


def test_waived_gap_renders_as_waived_not_required():
    g = SECTIONS[1]["gaps"][0]
    m = model([{"section": "team", "gap_id": g["id"], "reason": "founder-only company", "by": "admin"}])
    assert m.sections[1].gaps == [] and "waived by admin" in m.sections[1].waived[0]


def test_docx_has_markers_gaps_and_appendix():
    from docx import Document

    d = Document(io.BytesIO(ag.render_docx(model())))
    body = "\n".join(p.text for p in d.paragraphs)
    assert "[1]" in body and "[EVIDENCE REQUIRED: team members (name, role, bio)]" in body and "Not an offer." in body
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "opportunity:1" in cells and "tool:r:opportunity" in cells


def test_pptx_has_section_slides_and_appendix():
    from pptx import Presentation

    prs = Presentation(io.BytesIO(ag.render_pptx(model())))
    text_ = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Evidence appendix" in text_ and "[EVIDENCE REQUIRED" in text_ and "[2][1]" in text_


def test_html_escapes_and_marks():
    h = ag.render_html(
        ag.build_model(title="<script>x</script>", subtitle="", artefact="a", sections=SECTIONS, ref_labels=LABELS)
    )
    assert "<script>" not in h and "&lt;script&gt;" in h and "<sup>[1]</sup>" in h


def test_financial_model_uses_live_formulas():
    from openpyxl import load_workbook

    snaps = [{"period": f"2026-0{m}", "cash": 1_000_000 - m * 50_000, "revenue": 10_000 + m * 1000, "opex": 60_000, "net_burn": 50_000 - m * 1000,
              "ref": f"financial_snapshot:{m}"} for m in range(1, 7)]  # fmt: skip
    data = ag.financial_model_xlsx(snaps, currency="USD", horizon=24, trailing_months=3, min_cash_buffer=0, raise_amount=500_000,
                                   raise_month_offset=5, raise_probability=0.1, sources=[("financial_snapshot:1", "s", "x")], title="Model")  # fmt: skip
    wb = load_workbook(io.BytesIO(data))
    assert {"Inputs", "Assumptions", "Projection", "Runway", "Evidence"} <= set(wb.sheetnames)
    p = wb["Projection"]
    for c in ("C2", "D2", "E2", "F2", "G2", "G25"):
        assert str(p[c].value).startswith("="), c  # derived cells are formulas, never pasted values
    assert str(wb["Runway"]["B1"].value).startswith("=IFERROR(MATCH(")
    assert str(wb["Assumptions"]["B3"].value).startswith("=IFERROR((Inputs!C")
    assert wb["Inputs"]["B3"].value == 950_000, "sourced inputs are values"


def test_financial_model_without_inputs_flags_gaps():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(ag.financial_model_xlsx([], currency=None, horizon=12, trailing_months=3, min_cash_buffer=0,
                                                          raise_amount=None, raise_month_offset=None, raise_probability=None,
                                                          sources=[], title="Model")))  # fmt: skip
    assert "EVIDENCE REQUIRED" in str(wb["Inputs"]["A3"].value) and "EVIDENCE REQUIRED" in str(
        wb["Assumptions"]["B6"].value
    )


def test_budget_and_dd_checklist_formulas():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(ag.budget_xlsx([{"category": "Staff", "item": "Engineer", "amount": 90000},
                                                  {"category": "Equipment", "item": "Rig", "amount": 25000}], "USD", "Budget", [])))  # fmt: skip
    ws = wb["Budget"]
    assert any(str(c.value).startswith("=SUM(C3:C4)") for r in ws.iter_rows() for c in r)
    assert str(wb["By category"]["B2"].value).startswith("=SUMIF(")
    dd = load_workbook(
        io.BytesIO(ag.dd_checklist_xlsx([{"key": "cap_table", "title": "Cap table", "documents": []}], "DD"))
    )
    assert str(dd["DD checklist"]["D3"].value).startswith('=IF(B3>0,"Covered","[EVIDENCE REQUIRED]")')


# ----------------------------------------------------------------------------- composition
def test_compose_turns_missing_profile_inputs_into_gaps():
    me = {"id": "o1", "name": "Acme", "country": "US", "profile": {"stage": "seed"}}
    opp = {"id": "p1", "currency": "USD", "amount_min": 1, "factors": {}}
    r = ToolResult("opportunity", {}, facts=[{"text": "a", "kind": "fact", "evidence": ["opportunity:p1"], "basis": []},
                                             {"text": "b", "kind": "fact", "evidence": ["opportunity:p1"], "basis": []}])  # fmt: skip
    order = packages()["packages"]["vc_pitch"]["sections"]
    d = compose(
        order,
        me,
        opp,
        {"opportunity": r},
        lambda k: f"tool:run:{k}",
        {"budget": None, "esg": None, "recommendation": None},
    )
    assert any("team members" in g for g in d.gaps["team"]) and any("traction" in g for g in d.gaps["traction"])
    assert d.claims["company"][0]["evidence"] == ["organization:o1"] and "stage" in d.claims["company"][1]["text"]
    assert d.claims["opportunity"][0]["text"] == "a" and d.claims["ask"][0]["text"] == "b"


def test_all_eight_artefacts_are_covered_by_packages():
    arts = {a for p in packages()["packages"].values() for a in p["artefacts"]}
    assert arts == {"pitch_deck", "executive_summary", "investment_memo", "grant_narrative", "budget", "financial_model",
                    "technical_annex", "dd_checklist"}  # fmt: skip


# ----------------------------------------------------------------------------- compliance
def test_compliance_rules_cite_spans():
    items = [{"section": "ask", "index": 0, "text": "Returns are guaranteed and risk-free."},
             {"section": "company", "index": 1, "text": "We are the leading provider."},
             {"section": "financials", "index": 0, "text": "We will reach break-even."}]  # fmt: skip
    f = compliance_check.check_texts(items, has_disclaimer=True)
    rules = {x["rule"]: x for x in f}
    assert (
        rules["guarantee_of_returns"]["severity"] == "blocker"
        and rules["guarantee_of_returns"]["span"].lower() == "guaranteed"
    )
    assert "unsupported_superlative" in rules and "forward_looking_without_disclaimer" not in rules
    assert "forward_looking_without_disclaimer" in {
        x["rule"] for x in compliance_check.check_texts(items, has_disclaimer=False)
    }


async def test_compliance_review_blocks_on_blocker():
    out = await compliance_check.review([{"key": "ask", "claims": [{"text": "Guaranteed 20% return."}]}])
    assert out["status"] == "blocked" and out["blocking"]


# ----------------------------------------------------------------------------- retention + settings
def test_retention_policies_and_audit_immutability():
    assert duration("2y").days == 730 and duration("24h").total_seconds() == 86400
    p = policies({"agent_run": {"retain": "3y"}, "audit_log": {"retain": "1d"}})
    assert p["agent_run"]["retain"] == "3y" and p["audit_log"]["retain"] == "7y", "the audit log can't be overridden"


def test_settings_validation():
    validate("governance", {"allow_self_approval": True, "webhook_allowlist": ["https://hooks.example.org/"]})
    for bad in ({"allow_self_approval": "yes"}, {"unknown": 1}, {"webhook_allowlist": ["http://x"]}):
        with pytest.raises(Problem):
            validate("governance", bad)
    validate("llm_router", {"tiers": {"small": {"model": "m", "max_tokens": 1000}}})
    with pytest.raises(Problem):
        validate("llm_router", {"tiers": {"huge": {"model": "m"}}})
    with pytest.raises(Problem):
        validate("retention", {"audit_log": {"retain": "1d"}})
    with pytest.raises(Problem):
        validate("taxonomy", {"classes": {"crypto": {"label": "x"}}})
    validate("taxonomy", {"classes": {"grant": {"label": "Grants", "keywords": ["grant"]}}})


def test_extensions_are_interfaces_only():
    from cortex import extensions as ext

    status = ext.verify()
    assert set(status) == {"negotiation_assistant", "portfolio_optimizer", "scenario_planner", "finance_digital_twin",
                           "ecosystem_mapper", "cortex_federation"}  # fmt: skip
    assert all(v.startswith("off") for v in status.values()) and ext.get("portfolio_optimizer") is None
    with pytest.raises(ext.ExtensionMisconfigured):
        ext.register("portfolio_optimizer", object())


# ----------------------------------------------------------------------------- Graph adapters
def test_graph_mail_and_event_items():
    from cortex.l1_perception.adapters.msgraph import event_item, mail_item

    m = mail_item({"internetMessageId": "<m1>", "subject": "Hi", "from": {"emailAddress": {"address": "Ines@Q.example"}},
                   "toRecipients": [{"emailAddress": {"address": "me@inspironics.net"}}], "sentDateTime": "2026-09-01T10:00:00Z"}, "me@inspironics.net")  # fmt: skip
    assert m["external_id"] == "graph:<m1>" and m["from"] == "ines@q.example" and m["to"] == ["me@inspironics.net"]
    e = event_item({"iCalUId": "u1", "subject": "Call", "start": {"dateTime": "2026-09-01T10:00:00.0000000"},
                    "attendees": [{"emailAddress": {"address": "a@b.org"}}], "isCancelled": True})  # fmt: skip
    assert (
        e["status"] == "CANCELLED" and e["attendees"] == ["a@b.org"] and e["occurred_at"].startswith("2026-09-01T10:00")
    )


@respx.mock
async def test_graph_mail_adapter_fetches_with_client_credentials(monkeypatch):
    from cortex.l1_perception.adapters.base import FetchContext, get_adapter
    from cortex.l1_perception.registry import load_configs

    monkeypatch.setenv("MSGRAPH_CLIENT_SECRET", "s3cret")
    cfg = load_configs()["msgraph_mail"]
    respx.post(url__regex=r"https://login.microsoftonline.com/.*/token").mock(
        return_value=httpx.Response(200, json={"access_token": "t"})
    )
    route = respx.get(url__regex=r"https://graph.microsoft.com/v1.0/users/.*/messages.*").mock(
        return_value=httpx.Response(
            200,
            json={
                "value": [{"internetMessageId": "<m2>", "subject": "Re", "sentDateTime": datetime.now(UTC).isoformat()}]
            },
        )
    )
    ctx = FetchContext(min_interval_seconds=0, respect_robots=False)
    items = [i async for i in get_adapter("msgraph_mail").fetch(cfg, ctx)]
    await ctx.aclose()
    assert len(items) == 1 and route.calls[0].request.headers["Authorization"] == "Bearer t"


async def test_collateral_gaps_never_quote_stripped_claims():
    from cortex.l6_agency.orchestrator import _external_gaps
    from cortex.l7_governance.citation_checker import check_with_revisions, collateral_gaps

    class Nothing:
        async def resolve(self, ref):
            return None

        def in_scope(self, r):
            return True

    rep = await check_with_revisions(
        [{"text": "Aurora Capital committed USD 5,000,000.", "kind": "fact", "evidence": ["opportunity:x"]}],
        Nothing(),
        "t",
    )
    safe, stripped = collateral_gaps(rep)
    assert "Aurora" not in " ".join(safe) and stripped[0]["text"].startswith("Aurora")
    assert all("Aurora" not in g for g in _external_gaps(rep.gaps))
